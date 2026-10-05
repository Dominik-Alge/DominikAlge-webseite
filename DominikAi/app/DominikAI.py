import os
import glob
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
from google import genai
from google.genai import types
import openpyxl

load_dotenv()

app = FastAPI(title="DominikAI API")

# CORS restlos freigeben, um jegliche Browser-Blockaden bei Fehlern zu verhindern
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

GLOBAL_KNOWLEDGE_BASE = ""

def build_knowledge_index():
    """Liest die zentrale Index-Excel im metadata-Ordner ein"""
    global GLOBAL_KNOWLEDGE_BASE
    print("🚀 Starte ultrakompakte Indizierung via knowledge_index.xlsx...")
    
    excel_content = ""
    possible_paths = [
        "DominikAi/app/metadata/knowledge_index.xlsx",
        "app/metadata/knowledge_index.xlsx",
        "metadata/knowledge_index.xlsx"
    ]
    
    target_file = None
    for path in possible_paths:
        if os.path.exists(path):
            target_file = path
            break
            
    if not target_file:
        print("⚠️ WARNUNG: knowledge_index.xlsx wurde im metadata-Pfad nicht gefunden!")
        return

    try:
        wb = openpyxl.load_workbook(target_file, data_only=True)
        for sheet in wb.sheetnames:
            excel_content += f"\n[RUBRIK: {sheet}]\n"
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                row_text = " | ".join([str(cell).strip() for cell in row if cell is not None])
                if row_text.strip():
                    excel_content += row_text + "\n"
        
        GLOBAL_KNOWLEDGE_BASE = excel_content
        print(f"✅ Index erfolgreich geladen! {len(GLOBAL_KNOWLEDGE_BASE)} Zeichen im RAM.")

    except Exception as e:
        print(f"Fehler beim Einlesen der Excel-Matrix: {e}")

@app.on_event("startup")
async def startup_event():
    build_knowledge_index()

SYSTEM_PROMPT = """Du bist DominikAI, der persönliche, KI-gestützte Assistent auf der offiziellen Website von Dominik Alge (dominikalge.ch). 
Deine Aufgabe ist es, Fragen zu Dominiks Person, seinem Lebenslauf, seinen politischen Positionen, seiner Arbeit in der Industrie und seinen Publikationen kompetent, freundlich und präzise zu beantworten. 
Antworte immer in der Ich-Perspektive für Dominik oder als sein persönlicher Sprecher. Bleibe stets professionell und sachlich.

WICHTIGE SICHERHEITSREGEL: Antworte AUSSCHLIESSLICH auf Basis der unten bereitgestellten Informationen aus Dominiks Inhaltsverzeichnis/Wissensmatrix. Wenn eine Information nicht darin steht oder du die Antwort nicht weißt, erfinde auf keinen Fall etwas (keine Halluzinationen!). 
Antworte in diesem Fall stattdessen höflich auf Deutsch: „Dazu liegen mir aktuell leider keine genauen Informationen vor. Bitte wende dich bei spezifischen Fragen direkt an Dominik.“
"""

class ChatRequest(BaseModel):
    message: str

@app.post("/api/chat")
async def chat_endpoint(request: ChatRequest):
    try:
        # Falls die Excel beim Starten nicht geladen werden konnte, fangen wir das ab
        current_context = GLOBAL_KNOWLEDGE_BASE if GLOBAL_KNOWLEDGE_BASE else "Keine Daten geladen."
        
        # Wir fügen Prompt und Kontext sauber als lesbaren Text zusammen
        full_instructions = f"{SYSTEM_PROMPT}\n\nHIER IST DEINE STRUKTURIERTE WISSENSMATRIX:\n{current_context}"
        
        # Aufruf nach den neuesten Google-Vorgaben für gemini-3-flash-preview
        response = client.models.generate_content(
            model='gemini-3-flash-preview',
            contents=request.message,
            config=types.GenerateContentConfig(
                system_instruction=full_instructions,
                temperature=0.15
            )
        )
        return {"reply": response.text}
    except Exception as e:
        print(f"💥 Backend Absturz-Details: {e}")
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("DominikAI:app", host="0.0.0.0", port=10000, reload=True)

